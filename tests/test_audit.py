import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
import io

from audit import emailer
from audit.export import build_xlsx
from audit.privacy import scan_answers, scan_text
from audit.questions import SECTIONS, TABLES, iter_questions
from audit.storage import SqlStorage, StorageError, build_record

SAMPLES = sorted((Path(__file__).resolve().parents[1] / "samples").glob("sample_submission_*.json"))


def _sample(i=0):
    return json.loads(SAMPLES[i].read_text())


def test_question_ids_unique_and_tables_defined():
    ids = [q["id"] for _, q in iter_questions()]
    assert len(ids) == len(set(ids))
    for _, q in iter_questions():
        if q["type"] == "table":
            assert q["id"] in TABLES
    assert len(SECTIONS) == 9


@pytest.mark.parametrize("text", ["VAS 7 after training", "RPE: 8", "hamstring strain", "ACL",
                                  "player name: X"])
def test_privacy_flags_player_level_text(text):
    assert scan_text(text)


@pytest.mark.parametrize("text", ["VAS 0-10 scale", "CR-10 (0-10)", "RPE 6-20",
                                  "Availability status and expected return window", ""])
def test_privacy_allows_process_text(text):
    assert not scan_text(text)


def test_dashboard_top_limited_to_three():
    q = next(q for _, q in iter_questions() if q["id"] == "dashboard_top")
    assert q["max_selections"] == 3


def test_samples_are_process_only():
    for p in SAMPLES:
        assert scan_answers(json.loads(p.read_text())["answers"], {}) == []


def test_sql_storage_revisions(tmp_path, monkeypatch):
    monkeypatch.delenv("RENDER", raising=False)
    store = SqlStorage(f"sqlite:///{tmp_path / 'a.db'}")
    s = _sample()
    r1 = build_record(store, submission_id=None, submitter=s["submitter"],
                      submitter_role=s["submitter_role"], answers=s["answers"])
    store.save(r1)
    r2 = build_record(store, submission_id=r1["submission_id"], submitter=s["submitter"],
                      submitter_role=s["submitter_role"], answers=s["answers"])
    store.save(r2)
    assert (r1["revision"], r2["revision"]) == (1, 2)
    store.set_email_status(r2["submission_id"], 2, "sent to 1 recipient(s)")
    latest = store.list_latest()
    assert len(latest) == 1 and latest[0]["revision"] == 2
    assert latest[0]["email_status"] == "sent to 1 recipient(s)"
    assert store.get(r1["submission_id"], 1)["email_status"] == "pending"


def test_sqlite_refused_on_render(monkeypatch):
    monkeypatch.setenv("RENDER", "true")
    monkeypatch.delenv("AUDIT_DATABASE_URL", raising=False)
    with pytest.raises(StorageError):
        SqlStorage()


def test_xlsx_has_two_tabs_and_rows():
    s = _sample()
    rec = build_record(None, submission_id=None, submitter=s["submitter"],
                       submitter_role=s["submitter_role"], answers=s["answers"])
    wb = load_workbook(io.BytesIO(build_xlsx(rec, s["answers"])))
    assert wb.sheetnames == ["Report Inventory", "Audit Answers"]
    inv = wb["Report Inventory"]
    assert inv.max_row == 1 + len(s["answers"]["report_inventory"])
    answers_text = " ".join(str(c.value) for row in wb["Audit Answers"].iter_rows() for c in row)
    assert rec["submission_id"] in answers_text
    assert "Doctor's comment and recommendation" in answers_text


def test_xlsx_handles_partial_submission():
    s = _sample(1)
    rec = build_record(None, submission_id=None, submitter=s["submitter"],
                       submitter_role="", answers=s["answers"])
    wb = load_workbook(io.BytesIO(build_xlsx(rec, s["answers"])))
    text = " ".join(str(c.value) for row in wb["Audit Answers"].iter_rows() for c in row)
    assert "Other: Whiteboard in the medical room" in text


def test_email_missing_config_returns_status(monkeypatch):
    for k in ("AUDIT_EMAIL_TO", "SMTP_HOST", "SMTP_USER", "SMTP_FROM"):
        monkeypatch.delenv(k, raising=False)
    ok, status = emailer.send_xlsx({"submission_id": "X", "revision": 1, "submitter": "t",
                                    "submitted_at_utc": "now"}, b"x", "x.xlsx")
    assert not ok and status.startswith("not sent")


def test_email_failure_returns_status(monkeypatch):
    monkeypatch.setenv("AUDIT_EMAIL_TO", "a@example.invalid")
    monkeypatch.setenv("SMTP_HOST", "127.0.0.1")
    monkeypatch.setenv("SMTP_PORT", "1")  # nothing listens here
    monkeypatch.setenv("SMTP_FROM", "b@example.invalid")
    ok, status = emailer.send_xlsx({"submission_id": "X", "revision": 1, "submitter": "t",
                                    "submitted_at_utc": "now"}, b"x", "x.xlsx")
    assert not ok and status.startswith("failed")
