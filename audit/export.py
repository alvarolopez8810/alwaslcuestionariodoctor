"""XLSX export: two tabs, "Report Inventory" and "Audit Answers"."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .questions import OTHER, SECTIONS, TABLES

GREEN = "0B4C44"
GOLD = "DAB344"
HEADER_FILL = PatternFill("solid", fgColor=GREEN)
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="F3EBD0")
WRAP = Alignment(wrap_text=True, vertical="top")


def _clean_rows(table_id: str, rows) -> list[dict]:
    """Drop rows where every user-editable cell is blank."""
    keys = [k for k, _, kind, _ in TABLES[table_id]["columns"] if kind != "fixed"]
    return [r for r in (rows or []) if any(str(r.get(k) or "").strip() for k in keys)]


def answer_to_text(q: dict, value) -> str:
    """Flatten one answer to readable text for the spreadsheet."""
    t = q["type"]
    if value in (None, "", [], {}):
        return ""
    if t in ("text", "textarea"):
        return str(value).strip()
    if t == "number":
        return f"{value:g}" if isinstance(value, (int, float)) else str(value)
    if t == "radio":
        parts = [value.get("choice") or ""]
        if value.get("detail"):
            parts.append(f"Detail: {value['detail'].strip()}")
        return "\n".join(p for p in parts if p)
    if t == "multi":
        choices = [c for c in value.get("choices", []) if c != OTHER]
        if OTHER in value.get("choices", []):
            choices.append(f"Other: {value.get('other', '').strip() or '(not specified)'}")
        return "; ".join(choices)
    if t == "table":
        cols = TABLES[q["id"]]["columns"]
        lines = []
        for r in _clean_rows(q["id"], value):
            lines.append(" | ".join(f"{label}: {r.get(k) or '-'}" for k, label, _, _ in cols))
        return "\n".join(lines)
    return str(value)


def _style_header(ws, row: int, ncols: int):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP


def build_xlsx(record: dict, answers: dict) -> bytes:
    wb = Workbook()
    meta = [
        ("Submission ID", record["submission_id"]),
        ("Revision", record["revision"]),
        ("Submitter", record["submitter"]),
        ("Submitter role", record.get("submitter_role", "")),
        ("Submitted (UTC)", record["submitted_at_utc"]),
    ]

    # ---- Tab 1: Report Inventory ------------------------------------------
    ws = wb.active
    ws.title = "Report Inventory"
    cols = TABLES["report_inventory"]["columns"]
    headers = ["Submission ID", "Revision", "Row"] + [label for _, label, _, _ in cols]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for i, r in enumerate(_clean_rows("report_inventory", answers.get("report_inventory")), 1):
        ws.append([record["submission_id"], record["revision"], i]
                  + [r.get(k) or "" for k, _, _, _ in cols])
    widths = [20, 9, 6] + [26] * len(cols)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    # ---- Tab 2: Audit Answers ---------------------------------------------
    ws = wb.create_sheet("Audit Answers")
    for k, v in meta:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append([])
    ws.append(["Section", "Question ID", "Question", "Answer"])
    _style_header(ws, ws.max_row, 4)
    header_row = ws.max_row
    for section in SECTIONS:
        for q in section["questions"]:
            if q["id"] == "report_inventory":
                n = len(_clean_rows("report_inventory", answers.get("report_inventory")))
                text = f"{n} item(s) - see 'Report Inventory' tab"
            else:
                text = answer_to_text(q, answers.get(q["id"]))
            ws.append([section["title"], q["id"], q["label"], text or "(no answer)"])
    for col, w in zip("ABCD", (34, 22, 60, 80)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = WRAP
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_filename(record: dict) -> str:
    return f"{record['submission_id']}_r{record['revision']}_process_audit.xlsx"
